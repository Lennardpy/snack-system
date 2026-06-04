#!/usr/bin/env python3
"""Build a single Word document from the project's Markdown documentation.

Produces docs/SSS-Documentation.docx with a title page, section breaks per
document, and basic Markdown rendering (headings, lists, tables, code blocks,
bold/inline-code).
"""
import re
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

ROOT = Path(__file__).resolve().parent.parent
DOCS = ROOT / "docs"

# Order of the compiled document.
SECTIONS = [
    ("README.md", "Overview"),
    ("docs/ARCHITECTURE.md", "Architecture"),
    ("docs/DATA_MODEL.md", "Data Model"),
    ("docs/API.md", "API Reference"),
    ("docs/DEPLOYMENT.md", "Deployment"),
    ("docs/CARD_READER.md", "Card Reader & Kiosk Setup"),
    ("docs/CONGRESSUS.md", "Congressus Integration"),
    ("docs/ADMIN_GUIDE.md", "Admin Guide"),
    ("docs/SECURITY.md", "Security Model"),
]

ACCENT = RGBColor(0x1F, 0x6F, 0xEB)   # blue
CODE_BG = RGBColor(0x33, 0x33, 0x33)
MONO = "Consolas"

INLINE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def add_inline(paragraph, text):
    """Render **bold** and `code` spans inside a paragraph."""
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = MONO
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E)
        else:
            paragraph.add_run(part)


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(8)
    run = p.add_run("\n".join(lines))
    run.font.name = MONO
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x1A, 0x1A, 0x1A)
    # subtle shading
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), "F2F2F2")
    p._p.get_or_add_pPr().append(shd)


def add_table(doc, header, rows):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    for i, h in enumerate(header):
        cell = table.rows[0].cells[i]
        cell.paragraphs[0].clear()
        add_inline(cell.paragraphs[0], h)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            if i < len(cells):
                cells[i].paragraphs[0].clear()
                add_inline(cells[i].paragraphs[0], val)
    doc.add_paragraph()


def split_table_row(line):
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [c.strip() for c in line.split("|")]


def is_table_divider(line):
    return bool(re.match(r"^\s*\|?[\s:|-]+\|?\s*$", line)) and "-" in line


def render_markdown(doc, md, section_title):
    lines = md.split("\n")
    i = 0
    first_h1_skipped = False
    while i < len(lines):
        line = lines[i]

        # Code fence
        if line.strip().startswith("```"):
            block = []
            i += 1
            while i < len(lines) and not lines[i].strip().startswith("```"):
                block.append(lines[i])
                i += 1
            add_code_block(doc, block)
            i += 1
            continue

        # Table
        if line.strip().startswith("|") and i + 1 < len(lines) and is_table_divider(lines[i + 1]):
            header = split_table_row(line)
            rows = []
            i += 2
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_table_row(lines[i]))
                i += 1
            add_table(doc, header, rows)
            continue

        stripped = line.strip()

        # Headings
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            # Replace the document's own first H1 with the section title (already added)
            if level == 1 and not first_h1_skipped:
                first_h1_skipped = True
                i += 1
                continue
            heading = doc.add_heading(level=min(level, 4))
            heading.clear()
            add_inline(heading, text)
            i += 1
            continue

        # Horizontal rule
        if re.match(r"^([-*_])\1{2,}$", stripped):
            i += 1
            continue

        # Blockquote
        if stripped.startswith(">"):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            add_inline(p, stripped.lstrip("> ").strip())
            for run in p.runs:
                run.italic = True
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            i += 1
            continue

        # Bullet list
        m = re.match(r"^\s*[-*]\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            add_inline(p, m.group(1))
            i += 1
            continue

        # Numbered list
        m = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            add_inline(p, m.group(1))
            i += 1
            continue

        # Blank
        if not stripped:
            i += 1
            continue

        # Normal paragraph
        p = doc.add_paragraph()
        add_inline(p, stripped)
        i += 1


def main():
    doc = Document()

    # Base styles
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)

    # ── Title page ──
    doc.add_paragraph().add_run("\n\n\n")
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = t.add_run("SSS — Stress Snack System")
    run.bold = True
    run.font.size = Pt(30)
    run.font.color.rgb = ACCENT

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("System Documentation")
    r.font.size = Pt(16)
    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    desc = doc.add_paragraph()
    desc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    desc.add_run(
        "Self-service snack kiosk for the Stress student association.\n"
        "Next.js · PostgreSQL · Prisma · Railway · Congressus"
    ).font.size = Pt(11)

    # ── Table of contents (manual list) ──
    doc.add_page_break()
    toc = doc.add_heading("Contents", level=1)
    for idx, (_, title) in enumerate(SECTIONS, start=1):
        p = doc.add_paragraph(style="List Number")
        p.add_run(title)

    # ── Sections ──
    for path, title in SECTIONS:
        doc.add_page_break()
        h = doc.add_heading(title, level=1)
        md = (ROOT / path).read_text(encoding="utf-8")
        render_markdown(doc, md, title)

    out = DOCS / "SSS-Documentation.docx"
    doc.save(out)
    print(f"Wrote {out} ({out.stat().st_size // 1024} KB)")


if __name__ == "__main__":
    main()
